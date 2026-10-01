import csv
import argparse
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def le_csv(caminho):
    with open(caminho, encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def le_xlsx(caminho):
    ws = load_workbook(caminho, data_only=True).active
    linhas = list(ws.iter_rows(values_only=True))
    if not linhas:
        return []
    cabec = [str(c) if c is not None else f"col{i}" for i, c in enumerate(linhas[0])]
    return [
        {cabec[i]: linha[i] for i in range(len(cabec))}
        for linha in linhas[1:]
    ]


def filtra(linhas, campo, valor):
    if not campo or valor is None:
        return linhas
    return [l for l in linhas if str(l.get(campo, "")).strip() == str(valor).strip()]


def agrupa_e_soma(linhas, por, soma):
    total = {}
    for l in linhas:
        try:
            v = float(l.get(soma, 0) or 0)
        except (TypeError, ValueError):
            v = 0.0
        total[l.get(por, "?")] = total.get(l.get(por, "?"), 0.0) + v
    return total


def monta_resumo(total):
    geral = sum(total.values())
    resumo = []
    for chave, valor in sorted(total.items(), key=lambda x: -x[1]):
        pct = (valor / geral * 100) if geral else 0
        resumo.append({por: chave, "total": round(valor, 2), "pct": f"{pct:.1f}%"})
    resumo.append({por: "TOTAL", "total": round(geral, 2), "pct": "100%"})
    return resumo


def escreve_excel(resumo, saida):
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"

    campos = list(resumo[0].keys())
    ws.append(campos)

    fundo = PatternFill("solid", fgColor="1F4E78")
    fonte = Font(bold=True, color="FFFFFF")
    for i, _ in enumerate(campos, 1):
        c = ws.cell(row=1, column=i)
        c.fill = fundo
        c.font = fonte
        c.alignment = Alignment(horizontal="center")

    for linha in resumo:
        ws.append([linha[c] for c in campos])

    for i, campo in enumerate(campos, 1):
        largura = max(len(str(campo)), *(len(str(l[campo])) for l in resumo)) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(largura, 40)

    ws.freeze_panes = "A2"
    wb.save(saida)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("-i", "--input", required=True)
    ap.add_argument("-o", "--output", default="relatorio.xlsx")
    ap.add_argument("--filtro-campo")
    ap.add_argument("--filtro-valor")
    ap.add_argument("--agrupa", required=True)
    ap.add_argument("--soma", required=True)
    args = ap.parse_args()

    p = Path(args.input)
    if p.suffix.lower() == ".csv":
        linhas = le_csv(p)
    elif p.suffix.lower() in (".xlsx", ".xlsm"):
        linhas = le_xlsx(p)
    else:
        raise SystemExit("usa csv ou xlsx")

    print(f"{len(linhas)} linha(s) lidas")

    linhas = filtra(linhas, args.filtro_campo, args.filtro_valor)
    print(f"{len(linhas)} apos filtro")

    global por
    por = args.agrupa

    total = agrupa_e_soma(linhas, args.agrupa, args.soma)
    resumo = monta_resumo(total)
    escreve_excel(resumo, args.output)

    print(f"\nrelatorio: {args.output}\n")
    print(f"{'':<25} {'total':>12}  {'%':>6}")
    print("-" * 47)
    for l in resumo:
        print(f"{str(l[por])[:25]:<25} {l['total']:>12.2f}  {l['pct']:>6}")


if __name__ == "__main__":
    main()